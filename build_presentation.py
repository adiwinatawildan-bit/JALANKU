import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    # Set to 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6] # Blank layout

    # Color Palette Constants
    NAVY = RGBColor(15, 30, 60)         # #0F1E3C
    DARK_BLUE = RGBColor(24, 76, 120)    # #184C78
    LIGHT_BG = RGBColor(248, 250, 252)   # #F8FAFC
    WHITE = RGBColor(255, 255, 255)
    AMBER = RGBColor(217, 119, 6)        # #D97706
    EMERALD = RGBColor(16, 185, 129)     # #10B981
    SLATE = RGBColor(71, 85, 105)        # #475569
    DARK_GRAY = RGBColor(30, 41, 59)     # #1E293B
    CARD_BORDER = RGBColor(226, 232, 240)# #E2E8F0
    CARD_BG = RGBColor(255, 255, 255)
    ACCENT_BG = RGBColor(241, 245, 249)  # #F1F5F9

    def add_header(slide, title_text, category_text="SEMINAR KERJA PRAKTIK (KP) — SISTEM JALAN KU"):
        # Top banner background
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = NAVY
        top_bar.line.color.rgb = NAVY

        # Accent line
        acc_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.15), Inches(13.333), Inches(0.06))
        acc_bar.fill.solid()
        acc_bar.fill.fore_color.rgb = AMBER
        acc_bar.line.color.rgb = AMBER

        # Category text
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.7), Inches(0.3))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category_text.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = AMBER

        # Title text
        txBox_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.38), Inches(11.7), Inches(0.65))
        tf_title = txBox_title.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE

        # Footer
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.7), Inches(0.35))
        tf_foot = footer_box.text_frame
        p_foot = tf_foot.paragraphs[0]
        p_foot.text = "Sistem Informasi JALAN KU © 2026 | AI YOLOv8 & SPK TOPSIS"
        p_foot.font.size = Pt(9)
        p_foot.font.color.rgb = SLATE

    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.color.rgb = color

    def add_card(slide, left, top, width, height, title, body_bullets, bg_color=CARD_BG, border_color=CARD_BORDER, title_color=NAVY):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)

        tb = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.15), width - Inches(0.36), height - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(14)
        p0.font.bold = True
        p0.font.color.rgb = title_color
        p0.space_after = Pt(8)

        for b in body_bullets:
            p = tf.add_paragraph()
            p.text = "• " + b
            p.font.size = Pt(11)
            p.font.color.rgb = DARK_GRAY
            p.space_after = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 1: COVER
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s1, NAVY)

    # Accent decorative shape
    dec = s1.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Inches(9.5), Inches(0), Inches(3.833), Inches(7.5))
    dec.fill.solid()
    dec.fill.fore_color.rgb = RGBColor(24, 45, 85)
    dec.line.color.rgb = RGBColor(24, 45, 85)

    # Title Card Text
    tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(10.5), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "SEMINAR LAPORAN KERJA PRAKTIK"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = AMBER
    p.space_after = Pt(10)

    p2 = tf1.add_paragraph()
    p2.text = "SISTEM INFORMASI JALAN KU"
    p2.font.size = Pt(36)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_after = Pt(8)

    p3 = tf1.add_paragraph()
    p3.text = "Sistem Informasi Pelaporan dan Penentuan Prioritas Penanganan Kerusakan Jalan\nBerbasis Web Menggunakan AI YOLOv8 dan SPK TOPSIS"
    p3.font.size = Pt(16)
    p3.font.color.rgb = RGBColor(203, 213, 225)
    p3.space_after = Pt(35)

    p4 = tf1.add_paragraph()
    p4.text = "Disusun Oleh : Adiwinata Wildan (Mahasiswa Kerja Praktik)\nProgram Studi Teknik Informatika / Sistem Informasi\nTahun Akademik 2026"
    p4.font.size = Pt(12.5)
    p4.font.color.rgb = WHITE

    # -------------------------------------------------------------
    # SLIDE 2: LATAR BELAKANG
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s2, LIGHT_BG)
    add_header(s2, "Latar Belakang & Permasalahan")

    add_card(s2, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.1), 
             "1. Saluran Aduan Konvensional", 
             ["Pengaduan manual/media sosial sering tercecer dan tidak memiliki koordinat presisi.",
              "Dinas kesulitan memetakan sebaran titik kerusakan jalan secara sistematis.",
              "Masyarakat tidak mendapatkan kepastian dan transparansi tindak lanjut laporan."])

    add_card(s2, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.1), 
             "2. Penentuan Prioritas Subjektif", 
             ["Keterbatasan anggaran daerah menuntut penanganan jalan yang paling mendesak terlebih dahulu.",
              "Penilaian perbaikan sebelumnya sering bersifat subjektif tanpa parameter ilmiah terukur.",
              "Risiko kecelakaan fatal bagi pengendara motor kerap terabaikan jika tidak ada sistem pembobotan."])

    add_card(s2, Inches(8.8), Inches(1.6), Inches(3.6), Inches(5.1), 
             "3. Kebutuhan Inovasi Teknologi", 
             ["Diperlukan Computer Vision AI (YOLO) untuk memindai foto kerusakan secara otomatis.",
              "Diperlukan metode SPK TOPSIS untuk perangkingan prioritas penanganan yang adil dan objektif.",
              "Diperlukan sistem monitoring 3 fase (Before, Progress, After) yang transparan bagi publik."])

    # -------------------------------------------------------------
    # SLIDE 3: TUJUAN & MANFAAT SISTEM
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s3, LIGHT_BG)
    add_header(s3, "Tujuan Pengembangan & Manfaat Sistem")

    add_card(s3, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.1),
             "🎯 Tujuan Pengembangan Sistem",
             ["Membangun portal pengaduan kerusakan jalan berbasis geotagging GPS presisi (Leaflet/OSM).",
              "Mengintegrasikan model AI YOLOv8 untuk otomatisasi deteksi lubang (pothole), retakan (crack), dan longsor (landslide).",
              "Menerapkan algoritma SPK TOPSIS (4 Kriteria Baku) untuk merangking prioritas perbaikan jalan.",
              "Mempercepat alur penugasan dari Admin ke dinas teknis operasional (OPD Bina Marga).",
              "Menyediakan sistem pelacakan progres perbaikan jalan secara real-time dan transparan bagi publik."])

    add_card(s3, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.1),
             "💡 Manfaat Bagi Stakeholder",
             ["Masyarakat Umum: Kemudahan melapor secara online, pelacakan live status, dan evaluasi kepuasan (rating).",
              "Dinas Teknis (OPD): Penerimaan surat tugas terstruktur, kemudahan update progres fisik 0-100%, dan dokumentasi foto.",
              "Pemerintah Daerah (Admin): Pengambilan keputusan alokasi anggaran tepat sasaran dan berbasis data objektif.",
              "Akuntabilitas Publik: Ketersediaan peta sebaran titik rawan dan rekam jejak audit (Audit Trail)."])

    # -------------------------------------------------------------
    # SLIDE 4: ARSITEKTUR & STACK TEKNOLOGI
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s4, LIGHT_BG)
    add_header(s4, "Arsitektur Sistem & Teknologi yang Digunakan")

    add_card(s4, Inches(0.8), Inches(1.6), Inches(2.7), Inches(5.1),
             "🌐 Frontend",
             ["Framework Blade Template",
              "Tailwind CSS (Responsive)",
              "Leaflet.js & OpenStreetMap",
              "FontAwesome Icons",
              "Interaktif & Mobile Friendly"])

    add_card(s4, Inches(3.8), Inches(1.6), Inches(2.7), Inches(5.1),
             "⚙️ Backend",
             ["Framework Laravel 13",
              "PHP Version 8.3+ (v8.4)",
              "RESTful API Service",
              "Role-Based Access (RBAC)",
              "Supabase Storage / Local Disk"])

    add_card(s4, Inches(6.8), Inches(1.6), Inches(2.7), Inches(5.1),
             "🤖 AI Computer Vision",
             ["Python 3.12 Engine",
              "Ultralytics YOLOv8",
              "Model: model_terbaru_kaggle.pt",
              "Deteksi Pothole, Crack, Landslide",
              "Estimasi Luas Kerusakan (m²)"])

    add_card(s4, Inches(9.8), Inches(1.6), Inches(2.7), Inches(5.1),
             "🧮 Decision Support",
             ["Metode SPK TOPSIS",
              "Normalisasi Matriks Keputusan",
              "Solusi Ideal Positif & Negatif",
              "Skor Preferensi V (0.00 - 1.00)",
              "Ranking Otomatis TOP 10"])

    # -------------------------------------------------------------
    # SLIDE 5: HAK AKSES 4 ROLE PENGGUNA
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s5, LIGHT_BG)
    add_header(s5, "Matriks & Alur Kerja 4 Role Pengguna")

    add_card(s5, Inches(0.8), Inches(1.6), Inches(2.7), Inches(5.1),
             "1. Masyarakat",
             ["Login / Register",
              "Buat Laporan & Pin GPS",
              "Upload Foto Awal (1-3)",
              "Tracking Status Laporan",
              "Beri Rating & Feedback (1-5★)"],
             title_color=AMBER)

    add_card(s5, Inches(3.8), Inches(1.6), Inches(2.7), Inches(5.1),
             "2. Admin (Pengawas)",
             ["Validasi / Verifikasi Laporan",
              "Tolak Laporan Palsu",
              "Tandai Laporan Duplikat",
              "Jalankan AI YOLO & TOPSIS",
              "Disposisi Surat Tugas ke OPD"],
             title_color=DARK_BLUE)

    add_card(s5, Inches(6.8), Inches(1.6), Inches(2.7), Inches(5.1),
             "3. OPD Lapangan",
             ["Terima Surat Penugasan",
              "Mulai Survei Fisik Lokasi",
              "Input Progres Fisik (0-100%)",
              "Upload Foto Before/Progress/After",
              "Selesaikan Perbaikan Jalan"],
             title_color=EMERALD)

    add_card(s5, Inches(9.8), Inches(1.6), Inches(2.7), Inches(5.1),
             "4. Super Admin",
             ["Manajemen User & Akun",
              "Master Data OPD (Dinas)",
              "Pengaturan Bobot SPK TOPSIS",
              "Monitoring Audit Trail Log",
              "Pengaturan Konfigurasi Sistem"],
             title_color=NAVY)

    # -------------------------------------------------------------
    # SLIDE 6: FITUR AI YOLO
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s6, LIGHT_BG)
    add_header(s6, "Fitur Inovasi 1: Deteksi Citra Kerusakan Berbasis AI YOLOv8")

    add_card(s6, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.1),
             "🔍 Mekanisme Kerja Model AI YOLO",
             ["Foto laporan warga diproses oleh skrip engine ai_engine/yolo_detector.py.",
              "Model dilatih khusus untuk mendeteksi 3 kelas cacat jalan utama:",
              "  • Pothole (Lubang Jalan)",
              "  • Crack (Retakan Aspal)",
              "  • Landslide (Longsor / Jalan Ambles)",
              "Output inferensi menghasilkan bounding box presisi, tingkat keyakinan (confidence %), dan perkiraan luas area rusak (m²).",
              "Mekanisme pra-pemrosesan citra otomatis mengompresi gambar untuk efisiensi memori server."])

    add_card(s6, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.1),
             "📊 Konversi Otomatis ke Skala SPK",
             ["Hasil deteksi AI langsung mengisi parameter penilaian jalan:",
              "1. Longsor (Landslide):",
              "   -> C1 (Skala Kerusakan) = 4.6 - 5.0",
              "   -> C2 (Keselamatan) = 4.6 - 5.0 (Fatalitas Ekstrem)",
              "2. Lubang (Pothole):",
              "   -> >= 3 lubang: C1 = 4.1 | C2 = 4.2",
              "   -> 1-2 lubang: C1 = 3.5 - 3.8 | C2 = 3.7",
              "3. Retakan (Crack):",
              "   -> C1 = 2.4 - 3.2 | C2 = 2.4 - 2.8",
              "Menghilangkan unsur subjektivitas dalam pengukuran kondisi lapangan."])

    # -------------------------------------------------------------
    # SLIDE 7: FITUR SPK TOPSIS
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s7, LIGHT_BG)
    add_header(s7, "Fitur Inovasi 2: Sistem Pendukung Keputusan (SPK) TOPSIS")

    add_card(s7, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.1),
             "⚖️ 4 Kriteria Baku Penilaian (Total 100%)",
             ["C1 - Tingkat/Luas Kerusakan (Bobot: 40% | Benefit)",
              "   Mengukur dimensi fisik kerusakan dan estimasi luas (m²) via AI YOLO.",
              "C2 - Keselamatan Pengguna (Bobot: 25% | Benefit)",
              "   Mengukur potensi bahaya fatal kecelakaan bagi pengendara motor/mobil.",
              "C3 - Jumlah Laporan Tervalidasi (Bobot: 25% | Benefit)",
              "   Akumulasi aduan warga di ruas jalan yang sama (Crowdsourcing).",
              "C4 - Lama Belum Tertangani (Bobot: 10% | Benefit)",
              "   Jumlah hari sejak laporan masuk agar tidak ada laporan terbengkalai."])

    add_card(s7, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.1),
             "📈 Tahapan Perhitungan & Output Ranking",
             ["1. Pembentukan Matriks Keputusan (X) untuk seluruh laporan aktif.",
              "2. Normalisasi Matriks (R) menggunakan pembagi akar kuadrat.",
              "3. Pembobotan Matriks Ternormalisasi (Y = R * W).",
              "4. Penentuan Solusi Ideal Positif (A+) dan Negatif (A-).",
              "5. Perhitungan Jarak Euclidean (D+ dan D-).",
              "6. Skor Preferensi TOPSIS (V = D- / (D+ + D-)):",
              "   • Skor >= 0.65 -> SANGAT PRIORITAS (Merah)",
              "   • Skor >= 0.35 -> PRIORITAS TINGGI (Oranye)",
              "   • Skor >= 0.15 -> SEDANG (Kuning)",
              "   • Skor < 0.15  -> RENDAH (Hijau)"])

    # -------------------------------------------------------------
    # SLIDE 8: DOKUMENTASI 3 FASE & TRACKING
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s8, LIGHT_BG)
    add_header(s8, "Fitur Inovasi 3: Pemantauan Progres 3 Fase & Live Tracking")

    add_card(s8, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.1),
             "📷 1. Kondisi Awal (Before)",
             ["Foto diunggah pertama kali oleh masyarakat saat membuat laporan.",
              "Diverifikasi oleh tim Admin dan dianalisis oleh model AI YOLO.",
              "Menjadi acuan dasar tingkat keparahan awal sebelum ditangani dinas."])

    add_card(s8, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.1),
             "🔨 2. Pengerjaan (In Progress)",
             ["Diunggah secara bertahap oleh petugas lapangan OPD (misal: 30%, 70%).",
              "Mencatat tanggal pelaksanaan dan catatan teknis konstruksi di lapangan.",
              "Masyarakat dapat melihat bukti pengerjaan secara transparan."])

    add_card(s8, Inches(8.8), Inches(1.6), Inches(3.6), Inches(5.1),
             "✅ 3. Selesai (After)",
             ["Diunggah oleh dinas saat progres fisik mencapai 100%.",
              "Memperlihatkan kondisi aspal yang telah mulus dan tuntas diperbaiki.",
              "Membuka akses bagi pelapor untuk memberikan rating kepuasan (1-5★)."])

    # -------------------------------------------------------------
    # SLIDE 9: DEMONSTRASI ANTARMUKA (UI)
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s9, LIGHT_BG)
    add_header(s9, "Tampilan Antarmuka Sistem JALAN KU")

    add_card(s9, Inches(0.8), Inches(1.6), Inches(5.6), Inches(2.4),
             "📱 Portal Masyarakat & Form Pengaduan",
             ["Peta interaktif Leaflet dengan tombol 'Deteksi Lokasi GPS Saya'.",
              "Input deskripsi, patokan jalan, dan unggah multi-foto kondisi awal.",
              "Dashboard riwayat tiket dan linimasa status pelaporan real-time."])

    add_card(s9, Inches(6.8), Inches(1.6), Inches(5.6), Inches(2.4),
             "🖥️ Dashboard Operasional Admin & TOPSIS",
             ["Tabel TOP 10 Rekomendasi Prioritas Penanganan Jalan secara otomatis.",
              "Panel Verifikasi, Penolakan Laporan, dan Penandaan Duplikasi.",
              "Tombol instan untuk menjalankan analisis AI YOLO dan Hitung TOPSIS."])

    add_card(s9, Inches(0.8), Inches(4.2), Inches(5.6), Inches(2.5),
             "👷 Dashboard Penugasan OPD Lapangan",
             ["Daftar surat perintah tugas aktif yang didelegasikan Admin.",
              "Formulir pelaporan persentase progres fisik berkala (0% s/d 100%).",
              "Galeri unggah foto dokumentasi Before, In Progress, dan After."])

    add_card(s9, Inches(6.8), Inches(4.2), Inches(5.6), Inches(2.5),
             "⚙️ Panel Master Super Administrator",
             ["Manajemen Pengguna (User Management) & pengaturan role hak akses.",
              "Manajemen Master Data Instansi OPD dan wilayah kewenangan.",
              "Pengaturan bobot kriteria SPK TOPSIS (C1-C4) dan Audit Trail Log."])

    # -------------------------------------------------------------
    # SLIDE 10: HASIL & KESIMPULAN
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s10, LIGHT_BG)
    add_header(s10, "Hasil Implementasi & Kesimpulan Kerja Praktik")

    add_card(s10, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.1),
             "🏆 Hasil Implementasi Sistem",
             ["1. Berhasil mengintegrasikan pelaporan masyarakat berbasis peta interaktif dengan akurasi koordinat GPS.",
              "2. Model AI YOLOv8 mampu mendeteksi kerusakan jalan (pothole, crack, landslide) secara otomatis dari foto.",
              "3. Metode TOPSIS sukses menghasilkan perankingan prioritas penanganan jalan yang objektif dan transparan.",
              "4. Alur penugasan OPD hingga monitoring progres 3 fase berjalan terstruktur dan terdokumentasi rapi.",
              "5. Pengujian fungsional seluruh role (Masyarakat, Admin, OPD, Super Admin) berjalan sesuai kebutuhan (100% Valid)."])

    add_card(s10, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.1),
             "✨ Kesimpulan & Saran Pengembangan",
             ["Kesimpulan:",
              "Sistem Informasi JALAN KU berhasil menjadi solusi inovatif dalam modernisasi tata kelola infrastruktur jalan daerah yang akuntabel, cepat tanggap, dan partisipatif.",
              "Saran Pengembangan:",
              "1. Penambahan fitur push-notification via WhatsApp / Email kepada masyarakat pelapor.",
              "2. Pengembangan aplikasi berbasis mobile native (Android/iOS) untuk petugas lapangan saat offline di area pelosok.",
              "3. Penambahan integrasi sensor getaran (accelerometer) dari smartphone pengguna jalan."])

    # -------------------------------------------------------------
    # SLIDE 11: PENUTUP (Q&A)
    # -------------------------------------------------------------
    s11 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s11, NAVY)

    tb_end = s11.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.333), Inches(3.5))
    tf_end = tb_end.text_frame
    tf_end.word_wrap = True

    pe1 = tf_end.paragraphs[0]
    pe1.alignment = PP_ALIGN.CENTER
    pe1.text = "TERIMA KASIH"
    pe1.font.size = Pt(44)
    pe1.font.bold = True
    pe1.font.color.rgb = WHITE
    pe1.space_after = Pt(12)

    pe2 = tf_end.add_paragraph()
    pe2.alignment = PP_ALIGN.CENTER
    pe2.text = "SESI TANYA JAWAB (Q & A)\nSEMINAR LAPORAN KERJA PRAKTIK"
    pe2.font.size = Pt(20)
    pe2.font.bold = True
    pe2.font.color.rgb = AMBER
    pe2.space_after = Pt(24)

    pe3 = tf_end.add_paragraph()
    pe3.alignment = PP_ALIGN.CENTER
    pe3.text = "Sistem Informasi Pelaporan dan Penentuan Prioritas Penanganan Kerusakan Jalan (JALAN KU)\nPresenter: Adiwinata Wildan"
    pe3.font.size = Pt(13)
    pe3.font.color.rgb = RGBColor(203, 213, 225)

    # Save presentation
    output_path = os.path.join(os.getcwd(), "PRESENTASI_KP_JALAN_KU.pptx")
    prs.save(output_path)
    print(f"Presentation generated successfully at: {output_path}")

if __name__ == "__main__":
    create_presentation()

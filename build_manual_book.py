import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_manual_book():
    doc = Document()
    
    # Page setup - Standard A4 with 1 inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.27)  # A4
        section.page_height = Inches(11.69)
        
        # Configure Header & Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Manual Book — Sistem Informasi JALAN KU")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.italic = True
        hrun.font.color.rgb = RGBColor(140, 150, 160)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Sistem Informasi Pelaporan dan Penentuan Prioritas Penanganan Kerusakan Jalan | 2026")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(140, 150, 160)

    # Styles & Colors
    NAVY = RGBColor(15, 30, 60)
    DARK_BLUE = RGBColor(24, 76, 120)
    SLATE = RGBColor(71, 85, 105)
    DARK_GRAY = RGBColor(40, 40, 40)
    AMBER = RGBColor(217, 119, 6)

    def style_p(p, space_before=0, space_after=6, line_spacing=1.15):
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing

    def add_h1(text):
        p = doc.add_paragraph()
        style_p(p, space_before=16, space_after=8)
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(16)
        run.bold = True
        run.font.color.rgb = NAVY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        style_p(p, space_before=12, space_after=6)
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(13)
        run.bold = True
        run.font.color.rgb = DARK_BLUE
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        style_p(p, space_before=8, space_after=4)
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.bold = True
        run.font.color.rgb = NAVY
        return p

    def add_body(text, bold_prefix=None, italic=False):
        p = doc.add_paragraph()
        style_p(p, space_before=0, space_after=5, line_spacing=1.15)
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Calibri"
            r_pre.font.size = Pt(11)
            r_pre.bold = True
            r_pre.font.color.rgb = DARK_GRAY
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.italic = italic
        run.font.color.rgb = DARK_GRAY
        return p

    def add_bullet(text, bold_prefix=None, level=0):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Calibri"
            r_pre.font.size = Pt(10.5)
            r_pre.bold = True
            r_pre.font.color.rgb = DARK_GRAY
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(10.5)
        run.font.color.rgb = DARK_GRAY
        return p

    def set_cell_background(cell, hex_color):
        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading_elm)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'''
            <w:tcMar {nsdecls("w")}>
                <w:top w:w="{top}" w:type="dxa"/>
                <w:bottom w:w="{bottom}" w:type="dxa"/>
                <w:left w:w="{left}" w:type="dxa"/>
                <w:right w:w="{right}" w:type="dxa"/>
            </w:tcMar>
        ''')
        tcPr.append(tcMar)

    def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
                <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideV w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr.append(borders)

    def add_callout(text, title="CATATAN PENTING:"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.27)
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="0F2942"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        style_p(p, space_before=2, space_after=2, line_spacing=1.15)
        run_title = p.add_run(f"📌 {title} ")
        run_title.bold = True
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(10)
        run_title.font.color.rgb = NAVY
        
        run_text = p.add_run(text)
        run_text.font.name = "Calibri"
        run_text.font.size = Pt(10)
        run_text.font.color.rgb = DARK_GRAY
        
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_after = Pt(4)

    def add_screenshot_box(caption_text, desc_text=None):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.27)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=200, bottom=200, left=180, right=180)
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>
                <w:left w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>
                <w:bottom w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>
                <w:right w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_p(p, space_before=4, space_after=4)
        
        r1 = p.add_run("🖼️ [SCREENSHOT APLIKASI]\n")
        r1.font.name = "Calibri"
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(100, 116, 139)
        
        r2 = p.add_run(caption_text)
        r2.font.name = "Calibri"
        r2.font.size = Pt(9.5)
        r2.font.italic = True
        r2.font.color.rgb = RGBColor(148, 163, 184)
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(4)
        p_cap.paragraph_format.space_after = Pt(6)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = "Calibri"
        r_cap.font.size = Pt(9.5)
        r_cap.font.bold = True
        r_cap.font.color.rgb = SLATE
        
        if desc_text:
            p_desc = doc.add_paragraph()
            style_p(p_desc, space_before=0, space_after=6)
            r_desc = p_desc.add_run(f"Penjelasan Tampilan: {desc_text}")
            r_desc.font.name = "Calibri"
            r_desc.font.size = Pt(10)
            r_desc.font.italic = True
            r_desc.font.color.rgb = SLATE

    # ==========================================
    # 1. COVER PAGE
    # ==========================================
    p_cov_top = doc.add_paragraph()
    style_p(p_cov_top, space_before=72, space_after=12)
    p_cov_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    r_cov_tag = p_cov_top.add_run("PANDUAN PENGGUNA RESMI (USER MANUAL)")
    r_cov_tag.font.name = "Calibri"
    r_cov_tag.font.size = Pt(12)
    r_cov_tag.font.bold = True
    r_cov_tag.font.color.rgb = AMBER
    
    p_cov_title = doc.add_paragraph()
    style_p(p_cov_title, space_before=12, space_after=6)
    p_cov_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cov_title = p_cov_title.add_run("MANUAL BOOK\nSISTEM JALAN KU")
    r_cov_title.font.name = "Calibri"
    r_cov_title.font.size = Pt(28)
    r_cov_title.font.bold = True
    r_cov_title.font.color.rgb = NAVY
    
    p_cov_sub = doc.add_paragraph()
    style_p(p_cov_sub, space_before=6, space_after=40)
    p_cov_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cov_sub = p_cov_sub.add_run("Sistem Informasi Pelaporan dan Penentuan Prioritas Penanganan Kerusakan Jalan Berbasis Web")
    r_cov_sub.font.name = "Calibri"
    r_cov_sub.font.size = Pt(13)
    r_cov_sub.font.color.rgb = SLATE
    
    # Decorative Divider Line on Cover
    p_cov_div = doc.add_paragraph()
    p_cov_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cov_div = p_cov_div.add_run("—" * 35)
    r_cov_div.font.color.rgb = RGBColor(203, 213, 225)
    
    p_cov_meta = doc.add_paragraph()
    style_p(p_cov_meta, space_before=80, space_after=0)
    p_cov_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cov_meta = p_cov_meta.add_run("Versi Sistem: 1.0 (Produksi)\nKategori: Dokumentasi & Panduan Operasional Kerja Praktik\nTahun Rilis: 2026")
    r_cov_meta.font.name = "Calibri"
    r_cov_meta.font.size = Pt(10.5)
    r_cov_meta.font.color.rgb = DARK_GRAY
    
    doc.add_page_break()

    # ==========================================
    # 2. KATA PENGANTAR
    # ==========================================
    add_h1("KATA PENGANTAR")
    add_body("Puji dan syukur kami panjatkan ke hadirat Tuhan Yang Maha Esa atas berkat dan rahmat-Nya, penyusunan Buku Panduan Pengguna (Manual Book) untuk Sistem Informasi Pelaporan dan Penentuan Prioritas Penanganan Kerusakan Jalan (JALAN KU) dapat diselesaikan dengan baik.")
    add_body("Buku panduan ini disusun sebagai acuan operasional resmi bagi seluruh pengguna yang berinteraksi dengan sistem JALAN KU, mencakup masyarakat umum sebagai pelapor, tim verifikator/Admin di Command Center, dinas teknis operasional lapangan (OPD Bina Marga/PUPR), hingga Super Administrator yang bertanggung jawab atas pengelolaan master data dan tata kelola sistem.")
    add_body("Panduan ini memuat instruksi langkah demi langkah, penjelasan fungsi tombol dan formulir, navigasi antarmuka, serta prosedur troubleshooting kendala operasional. Diharapkan buku panduan ini dapat mempermudah pemanfaatan sistem secara optimal, meningkatkan responsivitas layanan publik, serta mendukung transparansi penanganan infrastruktur jalan.")
    add_body("Kami menyampaikan terima kasih yang sebesar-besarnya kepada seluruh pihak yang telah memberikan dukungan, masukan, dan dedikasi dalam perancangan hingga implementasi platform JALAN KU.")
    
    p_sign = doc.add_paragraph()
    style_p(p_sign, space_before=18, space_after=12)
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_sign = p_sign.add_run("Garut, September 2026\n\n\nTim Pengembang Sistem JALAN KU")
    r_sign.font.name = "Calibri"
    r_sign.font.size = Pt(11)
    r_sign.font.color.rgb = DARK_GRAY
    
    doc.add_page_break()

    # ==========================================
    # 3. DAFTAR ISI
    # ==========================================
    add_h1("DAFTAR ISI")
    
    toc_items = [
        ("KATA PENGANTAR", "ii"),
        ("DAFTAR ISI", "iii"),
        ("BAB I PENDAHULUAN", "1"),
        ("  1.1 Deskripsi Sistem", "1"),
        ("  1.2 Tujuan Sistem", "1"),
        ("  1.3 Ruang Lingkup", "2"),
        ("  1.4 Pengguna Sistem", "2"),
        ("BAB II KEBUTUHAN SISTEM", "3"),
        ("  2.1 Perangkat Keras (Hardware)", "3"),
        ("  2.2 Perangkat Lunak (Software)", "3"),
        ("  2.3 Kebutuhan Akses Sistem", "3"),
        ("BAB III AKSES SISTEM", "4"),
        ("  3.1 Halaman Awal", "4"),
        ("  3.2 Login", "4"),
        ("  3.3 Logout", "5"),
        ("  3.4 Akses Berdasarkan Role", "5"),
        ("BAB IV PANDUAN PENGGUNAAN SISTEM", "6"),
        ("  4.1 Panduan Role Masyarakat", "6"),
        ("    4.1.1 Halaman Dashboard Masyarakat", "6"),
        ("    4.1.2 Membuat Pengaduan Kerusakan Jalan", "6"),
        ("    4.1.3 Menentukan Lokasi & Geotagging Peta", "7"),
        ("    4.1.4 Mengunggah Foto Bukti Kerusakan", "7"),
        ("    4.1.5 Melihat Riwayat Laporan Saya", "8"),
        ("    4.1.6 Melacak Status & Detail Laporan", "8"),
        ("    4.1.7 Mengisi Ulasan & Feedback Kepuasan", "9"),
        ("    4.1.8 Mengelola Profil Akun", "9"),
        ("  4.2 Panduan Role Admin (Verifikator & Pengawas)", "10"),
        ("    4.2.1 Dashboard Admin & TOP 10 TOPSIS", "10"),
        ("    4.2.2 Manajemen Daftar Pengaduan Masuk", "10"),
        ("    4.2.3 Verifikasi, Penolakan, dan Duplikasi Laporan", "11"),
        ("    4.2.4 Menjalankan Analisis AI YOLO", "11"),
        ("    4.2.5 Menghitung Ulang Prioritas TOPSIS", "12"),
        ("    4.2.6 Menerbitkan Surat Penugasan ke OPD", "12"),
        ("    4.2.7 Pemantauan Log Audit Sistem", "13"),
        ("  4.3 Panduan Role OPD / Petugas Lapangan", "14"),
        ("    4.3.1 Dashboard Beban Kerja OPD", "14"),
        ("    4.3.2 Melihat Daftar Surat Tugas Masuk", "14"),
        ("    4.3.3 Memulai Survei Lapangan", "15"),
        ("    4.3.4 Input Progres Fisik & Upload Foto Pengerjaan", "15"),
        ("    4.3.5 Menyelesaikan Tugas Perbaikan", "16"),
        ("  4.4 Panduan Role Super Admin", "17"),
        ("    4.4.1 Dashboard Master Administrator", "17"),
        ("    4.4.2 Manajemen Pengguna (User Management)", "17"),
        ("    4.4.3 Manajemen Master Data OPD", "18"),
        ("    4.4.4 Manajemen Peran & Hak Akses (Roles)", "18"),
        ("    4.4.5 Pengaturan Kriteria & Bobot SPK TOPSIS", "19"),
        ("    4.4.6 Pemantauan Audit Trail Komprehensif", "19"),
        ("    4.4.7 Pengaturan Konfigurasi Sistem (Settings)", "20"),
        ("BAB V FITUR UTAMA SISTEM", "21"),
        ("  5.1 Pelaporan Kerusakan Jalan & Geotagging Presisi", "21"),
        ("  5.2 Peta Interaktif Sebaran Kerusakan (Leaflet Map)", "21"),
        ("  5.3 Deteksi Kerusakan Otomatis Berbasis AI YOLO", "22"),
        ("  5.4 SPK Penentuan Prioritas Metode TOPSIS (C1 - C4)", "22"),
        ("  5.5 Verifikasi, Penolakan, dan Deteksi Laporan Duplikat", "23"),
        ("  5.6 Manajemen Surat Perintah Penugasan OPD", "23"),
        ("  5.7 Monitoring Progres Lapangan Tiga Fase (Before, In Progress, After)", "24"),
        ("  5.8 Pelacakan Status Real-Time (Live Tracking)", "24"),
        ("  5.9 Portal Statistik & Data Terbuka Publik", "25"),
        ("  5.10 Survei Kepuasan Masyarakat (Rating & Feedback)", "25"),
        ("BAB VI PENANGANAN KENDALA (TROUBLESHOOTING)", "26"),
        ("  6.1 Gagal Login ke Sistem", "26"),
        ("  6.2 Lupa Kata Sandi (Password)", "26"),
        ("  6.3 Foto Gagal Diunggah", "27"),
        ("  6.4 Lokasi GPS / Peta Tidak Muncul", "27"),
        ("  6.5 Data Formulir Tidak Tersimpan", "27"),
        ("  6.6 Halaman Tidak Dapat Dibuka", "28"),
        ("  6.7 Proses AI YOLO Memerlukan Waktu", "28"),
        ("BAB VII PENUTUP", "29")
    ]
    
    tbl_toc = doc.add_table(rows=len(toc_items), cols=2)
    tbl_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_toc.autofit = False
    set_table_borders(tbl_toc, color="FFFFFF")  # No borders for TOC
    
    for idx, (title, page) in enumerate(toc_items):
        row = tbl_toc.rows[idx]
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(5.5)
        c1.width = Inches(0.77)
        
        p0 = c0.paragraphs[0]
        p1 = c1.paragraphs[0]
        style_p(p0, space_before=1, space_after=1)
        style_p(p1, space_before=1, space_after=1)
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
        is_bab = title.startswith("BAB") or title in ["KATA PENGANTAR", "DAFTAR ISI"]
        
        r0 = p0.add_run(title)
        r0.font.name = "Calibri"
        r0.font.size = Pt(10)
        r0.bold = is_bab
        r0.font.color.rgb = NAVY if is_bab else DARK_GRAY
        
        r1 = p1.add_run(page)
        r1.font.name = "Calibri"
        r1.font.size = Pt(10)
        r1.bold = is_bab
        r1.font.color.rgb = DARK_GRAY
        
    doc.add_page_break()

    # ==========================================
    # BAB I: PENDAHULUAN
    # ==========================================
    add_h1("BAB I PENDAHULUAN")
    
    add_h2("1.1 Deskripsi Sistem")
    add_body("Sistem Informasi JALAN KU adalah aplikasi berbasis web terintegrasi yang dirancang khusus untuk memfasilitasi pelaporan kerusakan jalan oleh masyarakat serta mengotomatisasi penentuan prioritas penanganan jalan bagi instansi dinas teknis terkait.")
    add_body("Aplikasi ini menggabungkan mekanisme partisipasi publik (crowdsourcing), teknologi Computer Vision berbasis kecerdasan buatan (Artificial Intelligence YOLOv8) untuk mengidentifikasi tingkat keparahan cacat jalan dari citra foto, serta metode Sistem Pendukung Keputusan (SPK) TOPSIS (Technique for Order Preference by Similarity to Ideal Solution) untuk merangking usulan penanganan ruas jalan secara ilmiah, objektif, dan transparan.")

    add_h2("1.2 Tujuan Sistem")
    add_body("Sistem JALAN KU dikembangkan dengan tujuan sebagai berikut:")
    add_bullet("Menyediakan saluran pengaduan kerusakan jalan yang mudah diakses publik dengan pencatatan koordinat geografis (geotagging) presisi.", bold_prefix="1. Aksesibilitas Publik: ")
    add_bullet("Mengotomatisasi identifikasi jenis kerusakan (lubang/pothole, retakan/crack, dan longsor/landslide) serta estimasi luasan area menggunakan model AI YOLO.", bold_prefix="2. Analisis Citra Otomatis: ")
    add_bullet("Menghasilkan urutan prioritas perbaikan jalan berbasis kriteria terukur (kondisi fisik, keselamatan pengguna, jumlah aduan, dan lama tunggu) melalui metode TOPSIS.", bold_prefix="3. Pengambilan Keputusan Objektif: ")
    add_bullet("Mempercepat alur koordinasi penerbitan surat tugas dari dinas pengawas (Admin) kepada dinas pelaksana teknis (OPD).", bold_prefix="4. Efisiensi Alur Kerja: ")
    add_bullet("Menyediakan rekam jejak dokumentasi pengerjaan 3 fase (Before, In Progress, After) yang dapat dipantau oleh masyarakat secara real-time.", bold_prefix="5. Transparansi & Akuntabilitas: ")

    add_h2("1.3 Ruang Lingkup")
    add_body("Ruang lingkup sistem informasi JALAN KU mencakup:")
    add_bullet("Modul Publik: Informasi portal terbuka, peta sebaran titik kerusakan interaktif, statistik penanganan, panduan cara kerja, dan pelacakan status laporan mandiri.")
    add_bullet("Modul Pelaporan Masyarakat: Autentikasi akun, formulir pembuatan aduan dengan GPS picker Leaflet/OpenStreetMap, riwayat aduan pribadi, dan pengisian rating kepuasan.")
    add_bullet("Modul Admin Verifikator: Validasi kelayakan laporan, penolakan aduan, penandaan duplikasi aduan berulang, pemicu komputasi AI YOLO, pemicu kalkulasi SPK TOPSIS, penerbitan surat tugas, dan monitoring audit log.")
    add_bullet("Modul OPD Lapangan: Penerimaan surat tugas, pembaruan status survei fisik, pelaporan progres bertahap (0% sampai 100%), unggah foto bukti kerja 3 fase, dan penyelesaian tugas.")
    add_bullet("Modul Super Admin: Tata kelola akun dan hak akses pengguna, master data instansi dinas (OPD), kustomisasi bobot persentase kriteria SPK, pengaturan toleransi duplikasi, dan audit log menyeluruh.")

    add_h2("1.4 Pengguna Sistem")
    add_body("Sistem JALAN KU memiliki 4 (empat) tingkatan hak akses pengguna (role) yang saling berinteraksi:")
    
    # Table of Roles
    tbl_r = doc.add_table(rows=5, cols=3)
    tbl_r.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_r.autofit = False
    set_table_borders(tbl_r)
    
    headers_r = ["Role Pengguna", "Kategori Akses", "Tugas & Wewenang Utama"]
    for i, h in enumerate(headers_r):
        cell = tbl_r.rows[0].cells[i]
        set_cell_background(cell, "0F1E3C")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        style_p(p, space_before=2, space_after=2)
        r = p.add_run(h)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    roles_data = [
        ("Masyarakat", "Publik / Pelapor", "Membuat aduan kerusakan jalan, menentukan pin koordinat GPS, mengunggah foto awal, memantau kemajuan pengerjaan, dan mengisi feedback kepuasan."),
        ("Admin", "Pengawas (Command Center)", "Memvalidasi laporan warga, mengeksekusi deteksi AI YOLO, menghitung urutan prioritas TOPSIS, dan menugaskan dinas teknis (OPD)."),
        ("OPD / Dinas", "Pelaksana Lapangan", "Menerima disposisi tugas, melaksanakan survei fisik lokasi, mengunggah progres berkala (0-100%) dan foto bukti kerja (Before/Progress/After)."),
        ("Super Admin", "Administrator Sistem", "Mengelola master akun pengguna, data dinas OPD, mengatur konfigurasi bobot parameter SPK TOPSIS, dan mengawasi seluruh log audit keamanan.")
    ]
    
    widths_r = [Inches(1.2), Inches(1.3), Inches(3.77)]
    for row_idx, data in enumerate(roles_data, start=1):
        row = tbl_r.rows[row_idx]
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            cell.width = widths_r[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            style_p(p, space_before=2, space_after=2)
            r = p.add_run(text)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.color.rgb = DARK_GRAY

    doc.add_page_break()

    # ==========================================
    # BAB II: KEBUTUHAN SISTEM
    # ==========================================
    add_h1("BAB II KEBUTUHAN SISTEM")
    
    add_h2("2.1 Perangkat Keras (Hardware)")
    add_body("Sistem JALAN KU dirancang responsif sehingga dapat diakses menggunakan berbagai perangkat komputer maupun perangkat bergerak (smartphone).")
    add_bullet("Perangkat Klien Masyarakat / Petugas Lapangan: Smartphone berbasis Android atau iOS yang dilengkapi modul kamera minimal 2 Megapiksel serta sensor GPS aktif, atau Komputer/Laptop dengan koneksi internet aktif.", bold_prefix="Klien Pengguna: ")
    add_bullet("Perangkat Klien Admin / Super Admin: Komputer PC atau Laptop dengan prosesor minimal Dual Core 2.0 GHz, RAM minimal 4 GB, dan resolusi layar minimal 1366 x 768 piksel untuk kenyamanan pemantauan dashboard.", bold_prefix="Klien Administrator: ")
    add_bullet("Spesifikasi Server Aplikasi: Server berbasis Linux/Windows dengan prosesor minimal 2 vCPU, RAM 4 GB (disarankan 8 GB untuk inferensi AI YOLO), ruang penyimpanan SSD minimal 20 GB, dan jaringan berkecepatan tinggi.", bold_prefix="Lingkungan Server: ")

    add_h2("2.2 Perangkat Lunak (Software)")
    add_body("Perangkat lunak yang dibutuhkan untuk menjalankan sistem meliputi:")
    add_bullet("Sistem Operasi Klien: Windows 10/11, macOS, Ubuntu Linux, Android 9.0 ke atas, atau iOS 13 ke atas.")
    add_bullet("Peramban Web (Browser): Google Chrome versi 90+, Mozilla Firefox versi 88+, Microsoft Edge versi 90+, atau Safari versi 14+ dengan fitur JavaScript dan Geolocation diaktifkan.")
    add_bullet("Arsitektur Backend: PHP versi 8.3+ (v8.4), Framework Laravel 13, Node.js versi 18+, Python versi 3.10+ (dengan library Ultralytics YOLOv8, OpenCV, Pillow, PyTorch), dan Basis Data PostgreSQL / MySQL.")

    add_h2("2.3 Kebutuhan Akses Sistem")
    add_body("Untuk menjamin kelancaran penggunaan sistem, pengguna wajib memastikan persyaratan akses berikut:")
    add_bullet("Koneksi Internet: Sambungan internet aktif dengan kecepatan minimal 1 Mbps untuk pengunggahan foto laporan dan rendering peta interaktif.", bold_prefix="Konektivitas: ")
    add_bullet("Izin Lokasi Peramban (Browser Geolocation Permission): Pengguna wajib memilih opsi 'Allow' atau 'Izinkan' saat peramban meminta izin akses lokasi perangkat untuk akurasi penentuan titik koordinat peta.", bold_prefix="Izin Geolocation: ")
    add_bullet("Izin Akses Kamera & Galeri: Diperlukan saat pengguna mengambil atau memilih foto dokumentasi kondisi jalan yang dilaporkan.", bold_prefix="Izin Media: ")

    doc.add_page_break()

    # ==========================================
    # BAB III: AKSES SISTEM
    # ==========================================
    add_h1("BAB III AKSES SISTEM")
    
    add_h2("3.1 Halaman Awal (Portal Publik)")
    add_body("Halaman Awal (Landing Page) merupakan beranda utama yang dapat diakses oleh publik tanpa perlu melakukan proses login terlebih dahulu. Halaman ini memuat navigasi menu publik, statistik agregat penanganan jalan di wilayah terkait, peta interaktif sebaran laporan, alur cara kerja pelaporan, dan tombol akses masuk/daftar.")
    add_body("Alur Membuka Halaman Awal:")
    add_bullet("Buka peramban web pada perangkat Anda (Google Chrome disarankan).")
    add_bullet("Ketikkan alamat URL sistem JALAN KU pada address bar (contoh: http://localhost:8000 atau domain resmi yang ditentukan).")
    add_bullet("Tekan tombol Enter. Sistem akan menampilkan Beranda Utama JALAN KU secara penuh.")

    add_screenshot_box("Gambar 3.1 Halaman Beranda Utama Publik JALAN KU", "Menampilkan bilah navigasi Beranda, Peta Kerusakan, Laporan Publik, Statistik, Cara Kerja, Tentang, serta tombol Login dan Register.")

    add_h2("3.2 Prosedur Login Pengguna")
    add_body("Proses login diperlukan bagi seluruh pengguna terdaftar (Masyarakat, Admin, OPD, dan Super Admin) untuk masuk ke panel kerja masing-masing sesuai wewenangnya.")
    add_body("Langkah-langkah Login:")
    add_bullet("Pada bilah navigasi kanan atas halaman Beranda, klik tombol 'Login'.")
    add_bullet("Sistem akan menampilkan Formulir Masuk ke Akun.")
    add_bullet("Masukkan alamat Email terdaftar pada kolom 'Email' (contoh: masyarakat@jalanku.go.id, admin@jalanku.go.id, opd@jalanku.go.id, atau superadmin@jalanku.go.id).")
    add_bullet("Masukkan Kata Sandi akun Anda pada kolom 'Password'.")
    add_bullet("Centang opsi 'Ingat Saya' apabila ingin menyimpan sesi autentikasi pada perangkat pribadi.")
    add_bullet("Klik tombol biru/navy bertuliskan 'Masuk ke Sistem'.")
    add_bullet("Sistem memverifikasi kredensial. Jika valid, pengguna otomatis diarahkan ke Dashboard sesuai role. Jika salah, muncul peringatan 'Kredensial yang diberikan tidak cocok dengan data kami.'")

    add_screenshot_box("Gambar 3.2 Halaman Formulir Login Autentikasi Pengguna", "Formulir input Email, Kata Sandi, tautan Lupa Kata Sandi, dan tombol Masuk ke Sistem.")

    add_h2("3.3 Prosedur Logout")
    add_body("Untuk menjaga keamanan akun, pengguna disarankan melakukan logout setelah selesai menggunakan aplikasi:")
    add_bullet("Pada pojok kanan atas bilah navigasi dashboard, klik menu dropdown Profil (menampilkan nama pengguna dan foto avatar).")
    add_bullet("Pilih dan klik opsi 'Logout' (ikon pintu keluar).")
    add_bullet("Sistem akan menghapus token sesi aktif dan mengembalikan pengguna ke halaman Beranda Publik dengan pesan notifikasi bahwa Anda telah berhasil keluar.")

    add_h2("3.4 Akses Berdasarkan Role")
    add_body("Sistem menerapkan kendali akses berbasis peran (Role-Based Access Control / RBAC). Setelah login berhasil, rute sistem akan mengarahkan pengguna ke ruang kerja masing-masing:")
    add_bullet("Role Masyarakat diarahkan ke rute: /masyarakat/dashboard", bold_prefix="Masyarakat: ")
    add_bullet("Role Admin diarahkan ke rute: /admin/dashboard", bold_prefix="Admin Verifikator: ")
    add_bullet("Role OPD Lapangan diarahkan ke rute: /opd/dashboard", bold_prefix="OPD Dinas Teknis: ")
    add_bullet("Role Super Admin diarahkan ke rute: /super-admin/dashboard", bold_prefix="Super Administrator: ")

    doc.add_page_break()

    # ==========================================
    # BAB IV: PANDUAN PENGGUNAAN SISTEM
    # ==========================================
    add_h1("BAB IV PANDUAN PENGGUNAAN SISTEM")

    # 4.1 MASYARAKAT
    add_h2("4.1 Panduan Role Masyarakat")
    add_body("Role Masyarakat diperuntukkan bagi warga yang ingin berpartisipasi melaporkan kondisi jalan rusak di wilayahnya, memantau transparansi tindak lanjut dinas, dan memberikan ulasan kepuasan.")

    add_h3("4.1.1 Halaman Dashboard Masyarakat")
    add_body("Halaman Dashboard Masyarakat menampilkan ringkasan data personal pelapor, meliputi:")
    add_bullet("Kartu Ringkasan Statistik: Total Laporan Diajukan, Laporan Terverifikasi, Sedang Dikerjakan, dan Laporan Selesai Ditangani.")
    add_bullet("Tombol Cepat '+ Buat Pengaduan Baru' di bagian atas dashboard.")
    add_bullet("Tabel Riwayat Laporan Terkini: Menampilkan 5 laporan terakhir lengkap dengan nomor tiket, nama jalan, status, tanggal pengajuan, dan tombol aksi detail.")

    add_screenshot_box("Gambar 4.1 Halaman Dashboard Utama Masyarakat", "Menampilkan kartu statistik laporan pribadi, tombol cepat Buat Pengaduan Baru, dan daftar ringkas laporan terkini.")

    add_h3("4.1.2 Membuat Pengaduan Kerusakan Jalan")
    add_body("Untuk membuat pengaduan baru, ikuti langkah-langkah berikut:")
    add_bullet("Dari bilah menu samping atau Dashboard, klik tombol '+ Buat Pengaduan Baru'. Sistem membuka rute /masyarakat/laporan/buat.")
    add_bullet("Bagian 1 — Informasi Pengaduan: Masukkan 'Judul Pengaduan' (contoh: Aspal Rusak Berlubang Parah di Dekat RSUD), 'Deskripsi Kerusakan & Kondisi Lapangan' (jelaskan perkiraan kedalaman lubang, luas area rusak, atau potensi bahaya), dan 'Informasi Tambahan / Patokan Lokasi' (contoh: Depan gerbang sekolah, 100m dari simpang tiga).")

    add_h3("4.1.3 Menentukan Lokasi & Geotagging Peta")
    add_body("Bagian 2 pada formulir pengaduan digunakan untuk menetapkan data geografis ruas jalan secara presisi:")
    add_bullet("Isi kolom teks 'Nama Ruas Jalan' (contoh: Jalan Cikajang KM 4.5), 'Kecamatan', 'Desa / Kelurahan', dan 'Detail Alamat'.")
    add_bullet("Untuk menentukan koordinat GPS: Klik tombol kuning bertuliskan 'Deteksi Lokasi GPS Saya'. Sistem akan meminta izin peramban dan secara otomatis menempatkan pin marker biru pada koordinat latitude & longitude Anda saat ini.")
    add_bullet("Jika ingin menyesuaikan posisi: Pengguna dapat menggeser (drag & drop) pin marker biru pada peta Leaflet atau mengklik langsung pada titik jalan yang dimaksud di peta.")

    add_h3("4.1.4 Mengunggah Foto Bukti Kerusakan")
    add_body("Bagian 3 formulir pengaduan mewajibkan pengunggahan foto kondisi awal:")
    add_bullet("Klik kotak area unggah foto atau seret berkas foto ke dalam area tersebut.")
    add_bullet("Pengguna dapat mengunggah minimal 1 foto dan maksimal 3 foto kondisi awal kerusakan (format JPG, JPEG, PNG dengan ukuran maksimal 5 MB per foto).")
    add_bullet("Setelah seluruh data lengkap, klik tombol hijau 'Kirim Pengaduan Sekarang'.")
    add_bullet("Hasil Akhir: Sistem menyimpan data ke basis data, membuat entri penomoran tiket unik resmi (contoh: #JLK-202609-WEZPT), menetapkan status awal 'DIAJUKAN', dan mengarahkan pengguna ke halaman detail laporan.")

    add_screenshot_box("Gambar 4.2 Formulir Pembuatan Pengaduan Kerusakan Jalan", "Formulir input judul, deskripsi, nama jalan, GPS picker peta interaktif Leaflet, dan area unggah multi-foto.")

    add_h3("4.1.5 Melihat Riwayat Laporan Saya")
    add_body("Menu 'Laporan Saya' (/masyarakat/laporan) menampilkan seluruh daftar pengaduan yang pernah dibuat oleh akun pelapor. Fitur pada halaman ini meliputi:")
    add_bullet("Penyaring Status (Filter): Menampilkan laporan berdasarkan status 'Semua', 'Diajukan', 'Diverifikasi', 'Ditugaskan', 'Sedang Dikerjakan', 'Selesai', atau 'Ditolak'.")
    add_bullet("Pencarian: Kolom pencarian berdasarkan nomor tiket atau nama ruas jalan.")
    add_bullet("Kartu/Tabel Laporan: Menampilkan foto thumbnail, nomor tiket, nama jalan, tanggal lapor, badge status berwarna, dan tombol 'Lihat Detail'.")

    add_h3("4.1.6 Melacak Status & Detail Laporan (Live Tracking)")
    add_body("Saat tombol 'Lihat Detail' diklik pada suatu laporan, halaman detail (/masyarakat/laporan/{id}) menampilkan informasi komprehensif:")
    add_bullet("Linimasa Progres (Tracking Stepper): Diagram alur interaktif yang menandai tahapan pengerjaan (Diajukan -> Diverifikasi -> Ditugaskan -> Sedang Disurvei -> Sedang Diperbaiki -> Selesai).")
    add_bullet("Hasil Deteksi AI YOLO: Menampilkan foto dengan kotak pembatas (bounding box) deteksi kerusakan fisik (pothole, crack, landslide) dan estimasi luas area.")
    add_bullet("Galeri Dokumentasi Lapangan: Menampilkan foto kondisi awal dari warga, foto survei lapangan dari OPD, foto saat pengerjaan berlangsung, dan foto setelah jalan selesai diperbaiki.")

    add_screenshot_box("Gambar 4.3 Halaman Detail & Live Tracking Progres Laporan", "Linimasa status penanganan, galeri foto dokumentasi 3 fase, informasi penugasan dinas, dan hasil analisis deteksi AI.")

    add_h3("4.1.7 Mengisi Ulasan & Feedback Kepuasan")
    add_body("Fitur ini aktif secara otomatis setelah laporan berstatus 'SELESAI':")
    add_bullet("Buka halaman detail laporan yang telah berstatus Selesai.")
    add_bullet("Pada kartu 'Ulasan & Kepuasan Masyarakat', pilih rating bintang (1 hingga 5 Bintang).")
    add_bullet("Tuliskan ulasan pengalaman pada kolom teks (contoh: Perbaikan sangat cepat dan aspal sangat mulus, terima kasih Bina Marga).")
    add_bullet("Klik tombol 'Kirim Penilaian'. Rating tersimpan permanen dan menjadi indikator evaluasi performa dinas.")

    add_h3("4.1.8 Mengelola Profil Akun")
    add_body("Melalui menu 'Profil' (/profil), masyarakat dapat memperbarui Nama Lengkap, Nomor Telepon, Alamat, serta mengganti Kata Sandi lama dengan Kata Sandi baru.")

    doc.add_page_break()

    # 4.2 ADMIN
    add_h2("4.2 Panduan Role Admin (Verifikator & Pengawas)")
    add_body("Role Admin bertindak sebagai unit pengendali operasional (Command Center) yang bertanggung jawab memvalidasi laporan warga, memicu analisis kecerdasan buatan, menentukan prioritas penanganan jalan, dan mendelegasikan surat tugas ke dinas teknis.")

    add_h3("4.2.1 Dashboard Admin & TOP 10 TOPSIS")
    add_body("Dashboard Admin (/admin/dashboard) memuat panel monitoring utama:")
    add_bullet("Metrik Status Laporan: Menampilkan jumlah tiket Diajukan (Menunggu Verifikasi), Diverifikasi, Ditugaskan ke OPD, Sedang Dikerjakan, Selesai, dan Ditolak.")
    add_bullet("Tabel TOP 10 Rekomendasi Prioritas Penanganan Jalan (TOPSIS): Menampilkan peringkat 10 ruas jalan paling mendesak berdasarkan skor preferensi TOPSIS (0.0000 - 1.0000), kategori prioritas (Sangat Prioritas, Prioritas Tinggi, Sedang, Rendah), tingkat kerusakan terdeteksi, dan alasan rekomendasi.")
    add_bullet("Tombol 'Hitung Ulang TOPSIS': Memicu kalkulasi ulang algoritma perangkingan secara instan.")

    add_screenshot_box("Gambar 4.4 Dashboard Operasional Admin & Tabel Prioritas TOPSIS", "Kartu ringkasan metrik status laporan, grafik distribusi, dan tabel TOP 10 Rekomendasi Prioritas TOPSIS.")

    add_h3("4.2.2 Manajemen Daftar Pengaduan Masuk")
    add_body("Menu 'Daftar Pengaduan' (/admin/laporan) memuat seluruh laporan masyarakat:")
    add_bullet("Filter Kategori Status: Tab penyaring laporan Diajukan, Diverifikasi, Ditugaskan, Proses, Selesai, dan Ditolak.")
    add_bullet("Tabel Laporan: Menampilkan nomor tiket, foto pratinjau, nama pelapor, nama jalan, kecamatan, status, tanggal, dan tombol aksi 'Periksa Laporan'.")

    add_h3("4.2.3 Verifikasi, Penolakan, dan Duplikasi Laporan")
    add_body("Pada halaman detail laporan (/admin/laporan/{id}), Admin memiliki wewenang eksekusi:")
    add_bullet("Verifikasi Laporan: Klik tombol hijau 'Verifikasi Laporan' jika laporan valid dan terbukti terdapat kerusakan jalan. Status berubah menjadi 'DIVERIFIKASI'.")
    add_bullet("Tolak Laporan: Klik tombol merah 'Tolak Laporan' jika laporan palsu, foto tidak relevan, atau bukan kewenangan. Masukkan alasan penolakan pada pop-up dialog (contoh: Foto tidak menampilkan jalan umum), lalu klik konfirmasi. Status berubah menjadi 'DITOLAK'.")
    add_bullet("Tandai Duplikat: Klik tombol kuning 'Tandai Duplikat' jika laporan melaporkan titik kerusakan yang sama dengan laporan sebelumnya. Sistem akan mengelompokkan laporan ke tiket induk dan memperbarui status menjadi 'DUPLIKAT'.")

    add_h3("4.2.4 Menjalankan Analisis AI YOLO")
    add_body("Admin dapat memicu pemindaian citra foto menggunakan model AI YOLOv8:")
    add_bullet("Pada halaman detail laporan, klik tombol biru bertuliskan 'Jalankan Analisis AI YOLO'.")
    add_bullet("Sistem secara asynchronous menjalankan skrip Python ai_engine/yolo_detector.py terhadap foto laporan.")
    add_bullet("Hasil: Sistem menampilkan deteksi objek (jumlah pothole, crack, landslide), confidence score (%), dan perkiraan luasan ($m^2$). Nilai ini secara otomatis memperbarui kriteria C1 (Skala Kerusakan) dan C2 (Keselamatan Pengguna) pada tabel penilaian jalan.")

    add_h3("4.2.5 Menghitung Ulang Prioritas TOPSIS")
    add_body("Untuk memperbarui seluruh ranking prioritas jalan setelah ada verifikasi baru:")
    add_bullet("Dari Dashboard Admin, klik tombol 'Hitung Ulang TOPSIS'.")
    add_bullet("Sistem membaca seluruh data laporan aktif, membentuk matriks keputusan ternormalisasi terbobot (C1-C4), menghitung jarak solusi ideal positif (D+) dan negatif (D-), serta menetapkan skor preferensi (V).")
    add_bullet("Daftar tabel TOP 10 langsung tersegarkan dengan urutan prioritas paling mutakhir.")

    add_h3("4.2.6 Menerbitkan Surat Penugasan ke OPD")
    add_body("Untuk menugaskan dinas teknis pelaksana:")
    add_bullet("Buka laporan berstatus 'DIVERIFIKASI'.")
    add_bullet("Klik tombol 'Tugaskan ke OPD'.")
    add_bullet("Pada jendela modal dialog, pilih instansi 'Dinas OPD Penanggung Jawab' (contoh: Dinas Bina Marga dan Penataan Ruang) dan ketikkan 'Instruksi / Catatan Khusus Penugasan' (contoh: Lakukan penambalan segera karena jalur padat kendaraan).")
    add_bullet("Klik 'Kirim Penugasan'. Status laporan otomatis diperbarui menjadi 'DITUGASKAN' dan muncul di dashboard OPD terkait.")

    add_screenshot_box("Gambar 4.5 Modal Dialog Penugasan Laporan ke Dinas Teknis (OPD)", "Pilihan instansi OPD pelaksana, kolom catatan instruksi kerja, dan tombol konfirmasi penugasan.")

    add_h3("4.2.7 Pemantauan Log Audit Sistem")
    add_body("Melalui menu 'Audit Logs' (/admin/audit-logs), Admin dapat meninjau rekaman aktivitas operasional: siapa yang melakukan verifikasi, kapan analisis AI dijalankan, dan kapan penugasan diterbitkan.")

    doc.add_page_break()

    # 4.3 OPD / PETUGAS
    add_h2("4.3 Panduan Role OPD / Petugas Lapangan")
    add_body("Role OPD (Organisasi Perangkat Daerah / Dinas Teknis) digunakan oleh tim operasional lapangan (Bina Marga / PUPR) untuk mengelola surat perintah tugas perbaikan jalan, melakukan survei fisik, memperbarui progres konstruksi, dan mengunggah dokumentasi foto pekerjaan.")

    add_h3("4.3.1 Dashboard Beban Kerja OPD")
    add_body("Dashboard OPD (/opd/dashboard) menyajikan ringkasan tugas instansi:")
    add_bullet("Kartu Metrik Tugas: Total Penugasan Masuk, Tugas Sedang Disurvei, Tugas Sedang Diperbaiki, dan Tugas Telah Selesai.")
    add_bullet("Tabel Penugasan Aktif: Menampilkan daftar perintah kerja terdekat lengkap dengan lokasi jalan, batas waktu, tingkat urgensi, dan status terkini.")

    add_screenshot_box("Gambar 4.6 Dashboard Penugasan Kerja Lapangan OPD", "Menampilkan kartu statistik penugasan dinas dan tabel surat perintah tugas aktif.")

    add_h3("4.3.2 Melihat Daftar Surat Tugas Masuk")
    add_body("Menu 'Daftar Tugas' (/opd/tugas) menampilkan seluruh surat perintah tugas yang didelegasikan oleh Admin kepada instansi OPD yang bersangkutan. Petugas dapat menyaring tugas berdasarkan status 'Ditugaskan', 'Survei', 'Progres', atau 'Selesai'.")

    add_h3("4.3.3 Memulai Survei Lapangan")
    add_body("Ketika tim lapangan diberangkatkan ke lokasi koordinat untuk melakukan peninjauan fisik dan pengukuran kebutuhan material:")
    add_bullet("Buka rincian tugas perbaikan terkait.")
    add_bullet("Klik tombol oranye 'Mulai Survei Lapangan'.")
    add_bullet("Hasil: Status penugasan dan status laporan warga diperbarui menjadi 'SEDANG DISURVEI' (SURVEI).")

    add_h3("4.3.4 Input Progres Fisik & Upload Foto Pengerjaan")
    add_body("Selama proses perbaikan jalan berlangsung, petugas wajib memperbarui kemajuan kerja secara berkala:")
    add_bullet("Pada halaman detail tugas, gulir ke panel 'Input Progres Fisik'.")
    add_bullet("Isi Formulir Progres: Masukkan 'Persentase Progres' (misal: 30 untuk 30%, 65 untuk 65%, 100 untuk selesai), tentukan 'Tanggal Pelaksanaan', 'Estimasi Tanggal Selesai', dan tuliskan 'Catatan Teknis Lapangan' (contoh: Pengerukan aspal rusak dan pemadatan agregat fondasi bawah).")
    add_bullet("Unggah Foto Dokumentasi: Pilih kategori foto yang diunggah: 'Sebelum Diperbaiki (Before)', 'Sedang Dikerjakan (In Progress)', atau 'Selesai (After)'. Lampirkan berkas foto dari kamera lapangan.")
    add_bullet("Klik tombol 'Simpan Progres'.")
    add_bullet("Hasil: Status laporan otomatis berubah menjadi 'SEDANG DIPERBAIKI'. Progres persentase dan foto dokumentasi langsung tampil di portal pelacakan publik dan masyarakat.")

    add_screenshot_box("Gambar 4.7 Formulir Pelaporan Progres Kerja Fisik & Upload Foto Tiga Fase", "Formulir persentase kemajuan 0-100%, catatan teknis, pemilihan kategori foto Before/Progress/After, dan tombol Simpan.")

    add_h3("4.3.5 Menyelesaikan Tugas Perbaikan")
    add_body("Tahap akhir penyelesaian pekerjaan perbaikan jalan:")
    add_bullet("Pastikan seluruh pekerjaan fisik telah selesai 100%.")
    add_bullet("Buka formulir progres, masukkan persentase '100%'.")
    add_bullet("Wajib mengunggah foto dokumentasi kategori 'Setelah Perbaikan (After)' yang memperlihatkan kondisi jalan yang telah mulus dan rapi.")
    add_bullet("Klik tombol hijau 'Selesaikan Tugas Pekerjaan'.")
    add_bullet("Hasil: Status tugas dan laporan resmi berubah menjadi 'SELESAI'. Sistem menutup penugasan dan secara otomatis membuka fitur pengisian rating/feedback bagi masyarakat pelapor.")

    doc.add_page_break()

    # 4.4 SUPER ADMIN
    add_h2("4.4 Panduan Role Super Admin")
    add_body("Role Super Admin memegang hak akses tertinggi atas tata kelola sistem, manajemen pengguna, konfigurasi master data instansi, pengaturan bobot parameter SPK, audit trail keamanan, dan preferensi platform.")

    add_h3("4.4.1 Dashboard Master Administrator")
    add_body("Dashboard Super Admin (/super-admin/dashboard) menampilkan ikhtisar menyeluruh:")
    add_bullet("Statistik Global Sistem: Total Pengguna Terdaftar, Total Instansi OPD Aktif, Total Laporan Masuk di Seluruh Wilayah, dan Status Layanan Server.")
    add_bullet("Grafik Tren Pengaduan Bulanan dan Sebaran Status Penanganan.")

    add_screenshot_box("Gambar 4.8 Dashboard Utama Super Administrator", "Ringkasan total akun pengguna, master data OPD, statistik global pengaduan, dan status modul AI.")

    add_h3("4.4.2 Manajemen Pengguna (User Management)")
    add_body("Menu 'Manajemen Pengguna' (/super-admin/users) digunakan untuk mengelola akun:")
    add_bullet("Menambah Pengguna Baru: Klik tombol '+ Tambah Pengguna'. Isi Nama Lengkap, Email, Nomor Telepon, Kata Sandi, dan tetapkan Role (Masyarakat, Admin, OPD, Super Admin). Jika memilih role OPD, pilih instansi dinasnya. Klik 'Simpan'.")
    add_bullet("Mengedit Pengguna: Klik tombol 'Edit' pada baris tabel pengguna untuk memperbarui data, mereset kata sandi, atau mengubah hak akses role.")
    add_bullet("Mengubah Status Akun: Mengaktifkan atau menonaktifkan akun pengguna (Blokir Akses).")
    add_bullet("Menghapus Pengguna: Menghapus akun yang tidak lagi aktif dengan konfirmasi keamanan.")

    add_h3("4.4.3 Manajemen Master Data OPD")
    add_body("Menu 'Master Data OPD' (/super-admin/opd) digunakan untuk mendata dinas teknis:")
    add_bullet("Menambah OPD Baru: Klik '+ Tambah OPD'. Masukkan Nama Dinas (contoh: Dinas Bina Marga dan Penataan Ruang), Nama Kepala Dinas, Email Kantor, Nomor Telepon Operasional, dan Alamat Kantor. Klik 'Simpan'.")
    add_bullet("Mengedit & Menghapus OPD: Mengubah kontak operasional dinas atau menghapus entri dinas.")

    add_h3("4.4.4 Manajemen Peran & Hak Akses (Roles)")
    add_body("Menu 'Manajemen Peran' (/super-admin/roles) memuat matriks hak akses wewenang sistem untuk masing-masing role guna memastikan keamanan data operasional.")

    add_h3("4.4.5 Pengaturan Kriteria & Bobot SPK TOPSIS")
    add_body("Menu 'Kriteria & Bobot SPK' (/super-admin/kriteria) memungkinkan Super Admin menyesuaikan bobot parameter pengambilan keputusan perbaikan jalan:")
    add_bullet("Daftar 4 Kriteria Baku: C1 (Tingkat/Luas Kerusakan - Benefit), C2 (Keselamatan Pengguna - Benefit), C3 (Jumlah Laporan Tervalidasi - Benefit), dan C4 (Lama Belum Tertangani - Benefit).")
    add_bullet("Mengubah Bobot: Masukkan nilai persentase baru pada kolom input (contoh default: C1=40%, C2=25%, C3=25%, C4=10%). Pastikan total akumulasi seluruh bobot bernilai tepat 100%.")
    add_bullet("Klik tombol 'Simpan Perubahan Bobot'. Sistem otomatis menerapkan bobot baru pada komputasi TOPSIS berikutnya.")

    add_screenshot_box("Gambar 4.9 Pengaturan Kriteria & Bobot Persentase SPK TOPSIS", "Tabel 4 kriteria baku C1-C4, tipe atribut Benefit, input persentase bobot (Total 100%), dan tombol Simpan.")

    add_h3("4.4.6 Pemantauan Audit Trail Komprehensif")
    add_body("Menu 'Audit Logs' (/super-admin/audit-logs) mencatat seluruh riwayat aktivitas di sistem: login pengguna, perubahan data, eksekusi AI YOLO, kalkulasi TOPSIS, modifikasi pengguna, hingga penghapusan data untuk kebutuhan audit keamanan.")

    add_h3("4.4.7 Pengaturan Konfigurasi Sistem (Settings)")
    add_body("Menu 'Pengaturan Sistem' (/super-admin/settings) mengelola konfigurasi teknis: nama aplikasi, logo platform, batas toleransi jarak duplikasi laporan (meter), pengaturan threshold AI YOLO, dan parameter notifikasi.")

    doc.add_page_break()

    # ==========================================
    # BAB V: FITUR UTAMA SISTEM
    # ==========================================
    add_h1("BAB V FITUR UTAMA SISTEM")
    
    fitur_list = [
        {
            "no": "5.1",
            "nama": "Pelaporan Kerusakan Jalan & Geotagging Presisi",
            "tujuan": "Memungkinkan masyarakat membuat laporan kerusakan jalan dengan titik koordinat GPS presisi dan bukti visual.",
            "akses": "Masyarakat -> Menu 'Buat Pengaduan' (/masyarakat/laporan/buat).",
            "langkah": "1. Isi judul dan deskripsi.\n2. Tentukan nama jalan dan wilayah.\n3. Klik 'Deteksi Lokasi GPS Saya' atau geser pin peta.\n4. Unggah 1-3 foto awal.\n5. Klik 'Kirim Pengaduan Sekarang'.",
            "hasil": "Laporan tersimpan di basis data dengan nomor tiket resmi (#JLK-xxxx) dan status awal 'DIAJUKAN'.",
            "catatan": "Izin lokasi pada peramban web wajib diaktifkan (Allow)."
        },
        {
            "no": "5.2",
            "nama": "Peta Interaktif Sebaran Kerusakan (Leaflet / OpenStreetMap)",
            "tujuan": "Menyajikan visualisasi spasial titik sebaran kerusakan jalan di seluruh wilayah secara terbuka dan transparan.",
            "akses": "Publik & Semua Pengguna -> Menu 'Peta' (/peta).",
            "langkah": "1. Buka menu Peta Kerusakan.\n2. Navigasi peta (zoom in/out, pan).\n3. Klik marker pin jalan untuk melihat pop-up ringkasan foto, nomor tiket, dan status penanganan.",
            "hasil": "Peta menampilkan sebaran pin berwarna sesuai status penanganan (Merah: Diajukan, Kuning: Diproses, Hijau: Selesai).",
            "catatan": "Data peta dimuat secara dinamis melalui REST API Geo-Reports."
        },
        {
            "no": "5.3",
            "nama": "Deteksi Kerusakan Otomatis Berbasis AI YOLOv8",
            "tujuan": "Mengotomatisasi identifikasi jenis cacat fisik (lubang/pothole, retakan/crack, longsor/landslide), menghitung jumlah titik, dan estimasi luas area ($m^2$) secara objektif.",
            "akses": "Admin -> Detail Laporan -> Tombol 'Jalankan Analisis AI YOLO'.",
            "langkah": "1. Admin membuka detail laporan terverifikasi.\n2. Klik tombol 'Jalankan Analisis AI YOLO'.\n3. Engine AI memproses citra secara asynchronous.",
            "hasil": "Citra beranotasi bounding box, tabel cacat terdeteksi, confidence score (%), dan pengisian otomatis nilai kriteria C1 dan C2.",
            "catatan": "Format foto yang didukung adalah JPG, JPEG, dan PNG."
        },
        {
            "no": "5.4",
            "nama": "Sistem Pendukung Keputusan (SPK) Metode TOPSIS",
            "tujuan": "Merangking prioritas perbaikan seluruh ruas jalan aktif secara matematis berdasarkan 4 kriteria terukur (C1 Luas Kerusakan 40%, C2 Keselamatan Pengguna 25%, C3 Jumlah Laporan 25%, C4 Lama Belum Tertangani 10%).",
            "akses": "Admin / Super Admin -> Dashboard Operasional -> Tombol 'Hitung Ulang TOPSIS'.",
            "langkah": "1. Klik tombol 'Hitung Ulang TOPSIS'.\n2. Algoritma menormalkan matriks keputusan dan menerapkan bobot kriteria.\n3. Menghitung jarak solusi ideal positif (D+) dan negatif (D-).\n4. Menghasilkan skor preferensi (V) antara 0.0000 - 1.0000.",
            "hasil": "Daftar ranking prioritas terbarui otomatis dengan kategori: Sangat Prioritas (V>=0.65), Prioritas Tinggi (V>=0.35), Sedang (V>=0.15), dan Rendah (V<0.15).",
            "catatan": "Laporan berstatus Ditolak atau Duplikat otomatis dieliminasi dari perhitungan."
        },
        {
            "no": "5.5",
            "nama": "Verifikasi, Penolakan, dan Deteksi Duplikasi Laporan",
            "tujuan": "Memastikan validitas data aduan masyarakat dan mengelompokkan laporan aduan berulang pada ruas jalan yang sama.",
            "akses": "Admin -> Menu 'Manajemen Laporan' -> Buka Detail Laporan.",
            "langkah": "1. Admin meninjau kesesuaian foto dan koordinat.\n2. Klik 'Verifikasi' untuk menyetujui.\n3. Klik 'Tolak Laporan' (dengan alasan tertulis) jika tidak valid.\n4. Klik 'Tandai Duplikat' jika laporan berulang.",
            "hasil": "Status laporan berubah secara resmi dan tercatat dalam log audit.",
            "catatan": "Laporan duplikat berkontribusi menaikkan nilai kriteria C3 (Crowdsourcing) pada tiket induk."
        },
        {
            "no": "5.6",
            "nama": "Manajemen Surat Perintah Penugasan OPD",
            "tujuan": "Mendelegasikan pelaksanaan perbaikan jalan yang telah terverifikasi kepada dinas teknis yang berwenang.",
            "akses": "Admin -> Detail Laporan -> Tombol 'Tugaskan ke OPD'.",
            "langkah": "1. Klik tombol 'Tugaskan ke OPD'.\n2. Pilih instansi dinas teknis tujuan.\n3. Masukkan catatan instruksi pengerjaan.\n4. Klik 'Kirim Penugasan'.",
            "hasil": "Status laporan menjadi 'DITUGASKAN' dan terbit surat perintah tugas pada akun OPD terkait.",
            "catatan": "Notifikasi penugasan tercatat dalam linimasa laporan."
        },
        {
            "no": "5.7",
            "nama": "Monitoring Progres Lapangan Tiga Fase (Before, In Progress, After)",
            "tujuan": "Menyediakan transparansi dan akuntabilitas pengerjaan fisik dinas di lapangan secara bertahap dari 0% hingga 100%.",
            "akses": "OPD -> Menu 'Daftar Tugas' -> Panel 'Input Progres Fisik'.",
            "langkah": "1. Masukkan persentase kemajuan (0-100%).\n2. Tentukan tanggal pelaksanaan dan estimasi selesai.\n3. Unggah foto lapangan kategori Before, In Progress, atau After.\n4. Klik 'Simpan Progres'.",
            "hasil": "Linimasa riwayat progres bertambah dan foto dokumentasi dapat dilihat oleh publik/pelapor.",
            "catatan": "Foto After wajib dilampirkan saat penyelesaian progres 100%."
        },
        {
            "no": "5.8",
            "nama": "Pelacakan Status Real-Time (Live Tracking)",
            "tujuan": "Memungkinkan pelapor dan masyarakat memantau tahapan tindak lanjut laporan secara langsung dan transparan.",
            "akses": "Masyarakat -> 'Laporan Saya' -> 'Lihat Detail', atau Publik -> 'Laporan Publik'.",
            "langkah": "1. Buka halaman detail laporan.\n2. Pantau diagram alur tahapan penanganan dan riwayat catatan dinas.",
            "hasil": "Pengguna mengetahui posisi tahapan laporan secara akurat.",
            "catatan": "Informasi diperbarui secara otomatis setiap ada pembaruan status."
        },
        {
            "no": "5.9",
            "nama": "Portal Statistik & Data Terbuka Publik",
            "tujuan": "Menyajikan data analitik penanganan infrastruktur jalan dalam bentuk diagram grafik dan angka statistik.",
            "akses": "Publik & Semua Pengguna -> Menu 'Statistik' (/statistik).",
            "langkah": "1. Buka menu Statistik pada navigasi publik.\n2. Tinjau grafik distribusi laporan per kecamatan, rasio penyelesaian, dan tren bulanan.",
            "hasil": "Visualisasi data analitik infrastruktur jalan yang informatif.",
            "catatan": "Data statistik diperbarui secara otomatis berdasarkan data aktual di basis data."
        },
        {
            "no": "5.10",
            "nama": "Survei Kepuasan Masyarakat (Rating & Feedback)",
            "tujuan": "Mengukur indeks kepuasan warga terhadap kecepatan dan mutu hasil perbaikan jalan oleh dinas.",
            "akses": "Masyarakat -> Detail Laporan (Status Selesai) -> Panel 'Ulasan'.",
            "langkah": "1. Pilih rating 1-5 bintang.\n2. Tuliskan komentar ulasan kepuasan.\n3. Klik 'Kirim Penilaian'.",
            "hasil": "Rating tersimpan dan menjadi indikator evaluasi kinerja instansi dinas.",
            "catatan": "Hanya dapat diisi 1 kali oleh akun masyarakat pembuat laporan."
        }
    ]

    for item in fitur_list:
        add_h2(f"{item['no']} Fitur {item['nama']}")
        add_body(item['tujuan'], bold_prefix="Tujuan / Fungsi: ")
        add_body(item['akses'], bold_prefix="Cara Mengakses: ")
        add_body(item['langkah'], bold_prefix="Langkah Penggunaan:\n")
        add_body(item['hasil'], bold_prefix="Hasil Akhir: ")
        add_body(item['catatan'], bold_prefix="Catatan Khusus: ")
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    doc.add_page_break()

    # ==========================================
    # BAB VI: PENANGANAN KENDALA (TROUBLESHOOTING)
    # ==========================================
    add_h1("BAB VI PENANGANAN KENDALA (TROUBLESHOOTING)")
    add_body("Bagian ini memuat panduan penanganan mandiri atas kendala teknis dan operasional yang mungkin terjadi selama penggunaan aplikasi JALAN KU:")

    troubleshoot_data = [
        ("6.1", "Gagal Login ke Sistem", 
         "1. Email atau kata sandi yang dimasukkan salah (periksa Caps Lock).\n2. Akun belum terdaftar di sistem.\n3. Status akun dinonaktifkan oleh Super Admin.",
         "1. Periksa kembali penulisan alamat email dan kata sandi.\n2. Jika belum memiliki akun, klik 'Daftar Akun Baru'.\n3. Hubungi Super Administrator jika akun Anda berstatus dinonaktifkan."),
        
        ("6.2", "Lupa Kata Sandi (Password)",
         "Pengguna lupa kombinasi kata sandi akun miliknya.",
         "1. Pada halaman Login, klik tautan 'Lupa Kata Sandi?'.\n2. Masukkan alamat email terdaftar dan ikuti instruksi reset sandi.\n3. Atau hubungi Super Admin untuk bantuan reset kata sandi melalui panel User Management."),
         
        ("6.3", "Foto Gagal Diunggah",
         "1. Ukuran berkas foto melebihi batas maksimal 5 MB.\n2. Format berkas bukan gambar yang didukung (selain JPG, JPEG, PNG).\n3. Koneksi internet terputus saat proses pengunggahan berlangsung.",
         "1. Kompres atau perkecil resolusi foto di bawah 5 MB sebelum diunggah.\n2. Pastikan ekstensi berkas berupa .jpg, .jpeg, atau .png.\n3. Pastikan koneksi internet stabil lalu coba unggah kembali."),
         
        ("6.4", "Titik Lokasi GPS / Peta Tidak Muncul",
         "1. Izin lokasi (Geolocation Permission) pada peramban web diblokir (Blocked).\n2. Sensor GPS pada smartphone dalam kondisi nonaktif.\n3. Layanan peta OpenStreetMap/Leaflet terhalang koneksi jaringan.",
         "1. Klik ikon gembok di sebelah kiri address bar peramban, ubah izin 'Location / Lokasi' menjadi 'Allow / Izinkan', lalu refresh halaman.\n2. Aktifkan GPS/Lokasi pada menu pengaturan perangkat.\n3. Pengguna tetap dapat menentukan lokasi secara manual dengan menggeser pin marker pada peta."),
         
        ("6.5", "Data Formulir Tidak Tersimpan",
         "1. Terdapat kolom isian wajib (bertanda bintang merah *) yang masih kosong.\n2. Sesi login telah berakhir (session timeout) karena terlalu lama tidak aktif.",
         "1. Periksa pesan peringatan berwarna merah di bawah kolom formulir dan lengkapi data yang diminta.\n2. Muat ulang (refresh) halaman, login kembali, dan lakukan pengisian formulir ulang."),
         
        ("6.6", "Halaman Tidak Dapat Dibuka",
         "1. Server aplikasi sedang dalam pemeliharaan (maintenance).\n2. Sambungan internet pengguna terputus.",
         "1. Periksa sambungan koneksi data internet atau Wi-Fi Anda.\n2. Coba buka kembali sistem setelah beberapa saat atau hubungi administrator jaringan."),
         
        ("6.7", "Proses AI YOLO Memerlukan Waktu",
         "Engine AI sedang memproses antrean citra beresolusi tinggi pada server.",
         "1. Harap menunggu sekitar 10 - 30 detik saat proses komputasi berlangsung.\n2. Jika proses terhenti, Admin dapat menekan kembali tombol 'Jalankan Analisis AI YOLO' pada halaman detail laporan.")
    ]

    tbl_tb = doc.add_table(rows=len(troubleshoot_data) + 1, cols=4)
    tbl_tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_tb.autofit = False
    set_table_borders(tbl_tb)
    
    headers_tb = ["No", "Kendala Teknis", "Kemungkinan Penyebab", "Solusi Penanganan"]
    for i, h in enumerate(headers_tb):
        cell = tbl_tb.rows[0].cells[i]
        set_cell_background(cell, "0F1E3C")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        style_p(p, space_before=2, space_after=2)
        r = p.add_run(h)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    widths_tb = [Inches(0.5), Inches(1.5), Inches(2.0), Inches(2.27)]
    for row_idx, data in enumerate(troubleshoot_data, start=1):
        row = tbl_tb.rows[row_idx]
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            cell.width = widths_tb[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            style_p(p, space_before=2, space_after=2)
            r = p.add_run(text)
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            r.font.color.rgb = DARK_GRAY

    doc.add_page_break()

    # ==========================================
    # BAB VII: PENUTUP
    # ==========================================
    add_h1("BAB VII PENUTUP")
    add_body("Buku Panduan Pengguna (Manual Book) ini disusun sebagai pedoman operasional lengkap bagi seluruh tingkatan pengguna dalam mengoperasikan Sistem Informasi Pelaporan dan Penentuan Prioritas Penanganan Kerusakan Jalan (JALAN KU).")
    add_body("Dengan adanya integrasi antara pelaporan partisipatif masyarakat, deteksi citra otomatis berbasis kecerdasan buatan (AI YOLOv8), sistem pendukung keputusan (SPK TOPSIS), serta alur kerja penugasan dan monitoring dinas teknis (OPD), sistem JALAN KU diharapkan mampu mewujudkan tata kelola pemeliharaan infrastruktur jalan yang responsif, transparan, objektif, dan akuntabel.")
    add_body("Seluruh pengguna diharapkan dapat menjalankan peran dan wewenangnya dengan baik sesuai panduan yang tertera di dalam dokumen ini demi terwujudnya infrastruktur jalan yang aman, nyaman, dan berkualitas bagi seluruh lapisan masyarakat.")
    
    add_callout(
        "Apabila terdapat pertanyaan teknis lebih lanjut, penambahan fitur, atau kendala sistem yang belum tercakup dalam manual book ini, silakan menghubungi Tim Administrator Sistem JALAN KU melalui saluran helpdesk resmi atau email: admin@jalanku.go.id.",
        title="INFORMASI BANTUAN & DUKUNGAN SISTEM:"
    )

    # Save document
    output_path = os.path.join(os.getcwd(), "MANUAL_BOOK_JALAN_KU.docx")
    doc.save(output_path)
    print(f"File Manual Book berhasil dibuat di: {output_path}")

if __name__ == "__main__":
    create_manual_book()
